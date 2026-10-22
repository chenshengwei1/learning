# -*- coding: utf-8 -*-
"""Generate HKT-PE BES FRS Document for dsagent Project"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from datetime import datetime

def set_cell_shading(cell, color):
    """Set cell background color"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def create_frs_document():
    """Create HKT-PE BES FRS document for dsagent"""
    
    # Create document
    doc = Document()
    
    # Set up styles
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    
    # ============ COVER PAGE ============
    for _ in range(3):
        doc.add_paragraph()
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('HKT-PE BES FRS')
    run.font.size = Pt(24)
    run.font.bold = True
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('dsagent Intelligent Assistant System')
    run.font.size = Pt(18)
    
    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = version.add_run('v1.0')
    run.font.size = Pt(16)
    
    for _ in range(5):
        doc.add_paragraph()
    
    # Document info table
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    info_data = [
        ('Issue:', 'V1.0'),
        ('Date:', datetime.now().strftime('%Y-%m-%d')),
        ('Prepared by:', 'System Analyst'),
        ('Approved by:', '<TBD>'),
    ]
    
    for i, (label, value) in enumerate(info_data):
        row = info_table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].width = Inches(1.5)
        row.cells[1].width = Inches(3)
    
    for _ in range(5):
        doc.add_paragraph()
    
    # Company info
    company = doc.add_paragraph()
    company.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = company.add_run('HUAWEI TECHNOLOGIES CO., LTD.')
    run.font.size = Pt(12)
    run.font.bold = True
    
    copyright = doc.add_paragraph()
    copyright.alignment = WD_ALIGN_PARAGRAPH.CENTER
    copyright.add_run('Copyright © 2024 Huawei Technologies Co., Ltd. All rights reserved.')
    
    doc.add_page_break()
    
    # ============ CHANGE HISTORY ============
    heading = doc.add_heading('Change History', level=1)
    
    change_table = doc.add_table(rows=2, cols=5)
    change_table.style = 'Table Grid'
    
    headers = ['Version', 'Date', 'Revised By', 'Section Modified', 'Description']
    for i, header in enumerate(headers):
        cell = change_table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, 'D9E1F2')
        cell.paragraphs[0].runs[0].font.bold = True
    
    change_data = [('0.1', '15/01/2024', 'System Analyst', 'All', 'Initial draft version.')]
    for i, row_data in enumerate(change_data):
        for j, cell_data in enumerate(row_data):
            change_table.rows[i+1].cells[j].text = cell_data
    
    doc.add_paragraph()
    doc.add_page_break()
    
    # ============ TABLE OF CONTENTS ============
    heading = doc.add_heading('Table of Contents', level=1)
    
    toc_items = [
        '1.1 CR Number',
        '1.2 Overview',
        '1.3 Business Rules and Assumptions',
        '1.4 Solution Design (CRM/GUI Layout)',
        '1.5 Solution Design (Integration Context)',
        '1.6 Solution Design (Interface Specification)',
        '1.7 Solution Design (CBS & MDS-C)',
        '1.8 Impact on the System',
        '1.9 Application Limitations',
        '1.10 References',
        '1.11 Appendix',
    ]
    
    for item in toc_items:
        p = doc.add_paragraph(item, style='List Number')
    
    doc.add_page_break()
    
    # ============ MAIN CONTENT ============
    heading = doc.add_heading('dsagent Intelligent Assistant System', level=1)
    
    # 1.1 CR Number
    doc.add_heading('1.1 CR Number', level=2)
    p = doc.add_paragraph()
    p.add_run('CR ID: ').bold = True
    p.add_run('PE_CR2024001')
    
    # 1.2 Overview
    doc.add_heading('1.2 Overview', level=2)
    doc.add_heading('1.2.1 Requirement Introduction', level=3)
    
    p = doc.add_paragraph()
    p.add_run('The dsagent (Data Science Agent) system is an intelligent assistant web application based on LangChain/LangGraph framework. It provides multi-turn conversation management, streaming AI responses, tool invocation capabilities, and context compression features.')
    
    p = doc.add_paragraph()
    p.add_run('Key Features:').bold = True
    
    features = [
        'Multi-turn conversation management with SQLite persistence',
        'Streaming AI responses via Server-Sent Events (SSE)',
        'Tool invocation system (file operations, Git, network requests)',
        'MCP (Model Context Protocol) tool integration',
        'Context compression and token management',
        'Web-based dashboard for session management',
    ]
    
    for feature in features:
        doc.add_paragraph(feature, style='List Bullet')
    
    # Subscription Matrix
    doc.add_heading('Subscription Matrix', level=4)
    
    sub_table = doc.add_table(rows=2, cols=4)
    sub_table.style = 'Table Grid'
    
    headers = ['Offer', 'Subscribe in SOP/CSP', 'Subscribe in ESP', 'Subscribe in dsagent Platform']
    for i, header in enumerate(headers):
        cell = sub_table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, 'D9E1F2')
        cell.paragraphs[0].runs[0].font.bold = True
    
    row_data = ['dsagent-Assistant', 'N/A', 'N/A', 'Yes']
    for i, data in enumerate(row_data):
        sub_table.rows[1].cells[i].text = data
    
    doc.add_paragraph()
    
    # 1.3 Business Rules
    doc.add_heading('1.3 Business Rules and Assumptions', level=2)
    
    rules_table = doc.add_table(rows=6, cols=4)
    rules_table.style = 'Table Grid'
    
    headers = ['Seq.', 'Rule Name', 'Rule Description', 'Remark']
    for i, header in enumerate(headers):
        cell = rules_table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, 'D9E1F2')
        cell.paragraphs[0].runs[0].font.bold = True
    
    rules_data = [
        ('1', 'Rule_001', 'Maximum context tokens: 1,000,000', ''),
        ('2', 'Rule_002', 'Recursion limit: 150 steps', ''),
        ('3', 'Rule_003', 'Tool call limit: 100 per session', ''),
        ('4', 'Rule_004', 'Input sanitization required', 'Anti-prompt injection'),
        ('5', 'Rule_005', 'Auto-compress at 70% token usage', 'Configurable threshold'),
    ]
    
    for i, row_data in enumerate(rules_data):
        for j, cell_data in enumerate(row_data):
            rules_table.rows[i+1].cells[j].text = cell_data
    
    doc.add_paragraph()
    
    # 1.4 Solution Design
    doc.add_heading('1.4 Solution Design (CRM/GUI Layout)', level=2)
    doc.add_heading('1.4.1 Solution Overview', level=3)
    
    p = doc.add_paragraph()
    p.add_run('The dsagent system provides a web-based UI with the following main screens:')
    
    screens = [
        'Chat Interface: Main conversation screen with message display and input',
        'Session Management: List, create, rename, and delete sessions',
        'Configuration: YAML-based graph configuration editor',
        'Database Manager: SQLite database browsing and management',
        'Log Manager: Real-time log level control',
        'Graph Visualizer: Mermaid.js based workflow visualization',
    ]
    
    for screen in screens:
        doc.add_paragraph(screen, style='List Bullet')
    
    # UI Element table
    doc.add_heading('UI Element Specification', level=4)
    
    element_table = doc.add_table(rows=4, cols=6)
    element_table.style = 'Table Grid'
    
    headers = ['Element', 'Type', 'Description', 'Validation Rule', 'Action', 'Remarks']
    for i, header in enumerate(headers):
        cell = element_table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, 'D9E1F2')
        cell.paragraphs[0].runs[0].font.bold = True
    
    elements = [
        ('Message Input', 'TextArea', 'User message input field', 'Max 2000 chars', 'Submit', ''),
        ('Send Button', 'Button', 'Trigger chat request', 'Enabled when input not empty', 'POST /api/chat/stream', ''),
        ('Session List', 'ListBox', 'Display all sessions', 'Auto-refresh on update', 'Select/Rename/Delete', ''),
    ]
    
    for i, row_data in enumerate(elements):
        for j, cell_data in enumerate(row_data):
            element_table.rows[i+1].cells[j].text = cell_data
    
    doc.add_paragraph()
    
    # 1.5 Integration Context
    doc.add_heading('1.5 Solution Design (Integration Context)', level=2)
    doc.add_heading('1.5.1 Chat Flow', level=3)
    
    flow_steps = [
        '1. User sends message via POST /api/chat/stream',
        '2. web_server.py sanitizes input and creates/gets session',
        '3. stream_agent() invokes LangGraph agent asynchronously',
        '4. Agent processes message through graph nodes: sanitize -> load_history -> retrieve -> build -> agent -> validate',
        '5. Conditional routing: tools loop or END',
        '6. Events streamed via SSE: tool_start, content, tool_end, done',
        '7. Session state persisted to SQLite checkpoint',
    ]
    
    for step in flow_steps:
        doc.add_paragraph(step, style='List Number')
    
    doc.add_heading('1.5.2 Tool Execution Flow', level=3)
    
    tool_steps = [
        '1. Agent decides to invoke tool based on user request',
        '2. execute_tools() node finds and invokes tool',
        '3. Tool result returned as ToolMessage',
        '4. Graph routes back to agent for next decision',
        '5. Max 8 consecutive tool calls before forced END',
    ]
    
    for step in tool_steps:
        doc.add_paragraph(step, style='List Number')
    
    doc.add_paragraph()
    
    # 1.6 Interface Specification
    doc.add_heading('1.6 Solution Design (Interface Specification)', level=2)
    doc.add_heading('1.6.1 REST API Interfaces', level=3)
    
    api_table = doc.add_table(rows=8, cols=3)
    api_table.style = 'Table Grid'
    
    headers = ['Interface', 'Method', 'Description']
    for i, header in enumerate(headers):
        cell = api_table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, 'D9E1F2')
        cell.paragraphs[0].runs[0].font.bold = True
    
    apis = [
        ('POST /api/chat/stream', 'SSE', 'Streaming chat interface'),
        ('POST /api/chat', 'JSON', 'Non-streaming chat'),
        ('GET /api/sessions', 'JSON', 'List all sessions'),
        ('POST /api/sessions/new', 'JSON', 'Create new session'),
        ('GET /api/sessions/{id}/messages', 'JSON', 'Get session messages'),
        ('DELETE /api/sessions/{id}', 'JSON', 'Delete session'),
        ('GET /api/graph', 'HTML', 'Graph visualization'),
    ]
    
    for i, row_data in enumerate(apis):
        for j, cell_data in enumerate(row_data):
            api_table.rows[i+1].cells[j].text = cell_data
    
    doc.add_heading('1.6.2 MCP Tool Interfaces', level=3)
    
    p = doc.add_paragraph()
    p.add_run('MCP servers configured:')
    
    mcp_list = [
        'github: GitHub Copilot MCP (streamable_http)',
        'local_api: Local API MCP (stdio)',
        'local_http: HTTP MCP endpoint',
    ]
    
    for mcp in mcp_list:
        doc.add_paragraph(mcp, style='List Bullet')
    
    doc.add_paragraph()
    
    # 1.7 CBS & MDS-C
    doc.add_heading('1.7 Solution Design (CBS & MDS-C)', level=2)
    p = doc.add_paragraph()
    p.add_run('N/A - This is a general-purpose AI assistant, not a telecom billing system.')
    
    # 1.8 Impact
    doc.add_heading('1.8 Impact on the System', level=2)
    
    impacts = [
        'Database: SQLite agent_memory.db for session persistence',
        'API: DeepSeek API for LLM inference',
        'MCP: External tool servers (GitHub, local APIs)',
        'File System: Project file access via tools',
    ]
    
    for impact in impacts:
        doc.add_paragraph(impact, style='List Bullet')
    
    doc.add_paragraph()
    
    # 1.9 Limitations
    doc.add_heading('1.9 Application Limitations', level=2)
    
    limitations = [
        '1. Single-threaded event loop for session processing',
        '2. SQLite concurrent write limitations',
        '3. Tool execution timeout: 30 seconds',
        '4. No persistent user authentication',
        '5. Out Of Scope in this CR as requirement is not confirmed yet: Multi-model support',
        '6. Out Of Scope in this CR as requirement is not confirmed yet: Distributed deployment',
    ]
    
    for limit in limitations:
        doc.add_paragraph(limit, style='List Number')
    
    doc.add_paragraph()
    
    # 1.10 References
    doc.add_heading('1.10 References', level=2)
    
    references = [
        '<< LangChain Documentation >>',
        '<< LangGraph Documentation >>',
        '<< DeepSeek API Reference >>',
        '<< MCP Protocol Specification >>',
    ]
    
    for ref in references:
        doc.add_paragraph(ref, style='List Bullet')
    
    doc.add_paragraph()
    
    # 1.11 Appendix
    doc.add_heading('1.11 Appendix', level=2)
    doc.add_heading('1.11.1 Data Model', level=3)
    
    p = doc.add_paragraph()
    p.add_run('AgentState structure:').bold = True
    
    state_fields = [
        'messages: Annotated[list, operator.add]',
        'session_id: str',
        'knowledge: List[dict]',
        'history: List[dict]',
        'validation: str',
        'error_msg: str',
        'retry_count: int',
        'full_messages: List[Any]',
        'context_messages: List[Any]',
        'compression_info: dict',
        'llm_response: Any',
        'llm_start_time: float',
        'llm_duration_ms: float',
        'updated_context: List[Any]',
        'thread_id: str',
    ]
    
    for field in state_fields:
        doc.add_paragraph(field, style='List Bullet')
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Key Configuration:').bold = True
    
    config_items = [
        'MAX_CONTEXT_TOKENS: 1,000,000',
        'WARNING_THRESHOLD: 70%',
        'CAUTION_THRESHOLD: 85%',
        'DANGER_THRESHOLD: 95%',
        'AGENT_RECURSION_LIMIT: 150 (env)',
    ]
    
    for item in config_items:
        doc.add_paragraph(item, style='List Bullet')
    
    # Save document
    output_path = 'HKT-PE BES FRS dsagent v1.0.docx'
    doc.save(output_path)
    print(f'Document saved: {output_path}')
    return output_path

if __name__ == '__main__':
    create_frs_document()

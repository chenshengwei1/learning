#!/usr/bin/env python3
"""写入或追加数据到 Excel 文件。"""
import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill


def parse_data(data_arg: str) -> pd.DataFrame:
    p = Path(data_arg)
    if p.exists():
        raw = json.loads(p.read_text(encoding="utf-8"))
    else:
        raw = json.loads(data_arg)
    # 支持 [{"a":1},{"a":2}] 或 {"a":[1,2]}
    if isinstance(raw, list):
        return pd.DataFrame(raw)
    if isinstance(raw, dict):
        return pd.DataFrame(raw)
    raise ValueError("data 必须是 JSON 数组或对象")


def style_sheet(path: str, sheet: str):
    wb = load_workbook(path)
    ws = wb[sheet]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)
    wb.save(path)


def main():
    p = argparse.ArgumentParser(description="Write Excel")
    p.add_argument("--file", required=True)
    p.add_argument("--sheet", default="Sheet1")
    p.add_argument("--data", required=True, help="JSON 字符串或 JSON 文件路径")
    p.add_argument("--mode", choices=["write", "append"], default="write")
    p.add_argument("--header", default="true")
    p.add_argument("--no-style", action="store_true")
    args = p.parse_args()

    header = args.header.lower() == "true"
    try:
        df = parse_data(args.data)
    except Exception as e:
        print(json.dumps({"error": f"data parse failed: {e}"}, ensure_ascii=False))
        sys.exit(1)

    path = Path(args.file)

    if args.mode == "append" and path.exists():
        try:
            old = pd.read_excel(path, sheet_name=args.sheet)
            df = pd.concat([old, df], ignore_index=True)
        except ValueError:
            pass  # sheet 不存在，直接新建

    try:
        with pd.ExcelWriter(path, engine="openpyxl", mode="a" if path.exists() and args.mode == "append" else "w",
                            if_sheet_exists="replace" if path.exists() and args.mode == "append" else None) as writer:
            df.to_excel(writer, sheet_name=args.sheet, index=False, header=header)
    except Exception as e:
        print(json.dumps({"error": f"write failed: {e}"}, ensure_ascii=False))
        sys.exit(1)

    if not args.no_style and header:
        try:
            style_sheet(str(path), args.sheet)
        except Exception:
            pass

    print(json.dumps({
        "saved": str(path),
        "sheet": args.sheet,
        "rows": len(df),
        "mode": args.mode,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()